import ast
import csv
import io
from datetime import datetime
from typing import BinaryIO

from openpyxl import load_workbook

from app.schemas.post import PostCreate


def parse_posts(file: BinaryIO, extension: str) -> list[PostCreate]:
    """Читает .xlsx или .csv и превращает строки в список постов."""
    if extension == "csv":
        return parse_csv(file)

    return parse_excel(file)


def parse_excel(file: BinaryIO) -> list[PostCreate]:
    """Читает .xlsx-файл и превращает строки листа в список постов."""
    workbook = load_workbook(file, read_only=True)
    sheet = workbook.active

    rows = sheet.iter_rows(values_only=True)

    headers = next(rows)

    return [make_post(dict(zip(headers, row))) for row in rows]


def parse_csv(file: BinaryIO) -> list[PostCreate]:
    """Читает .csv-файл и превращает строки в список постов."""
    text = read_text(file).replace("\r\n", "\n")
    reader = csv.DictReader(io.StringIO(text), delimiter=detect_delimiter(text))

    return [make_post(row) for row in reader]


def read_text(file: BinaryIO) -> str:
    """Декодирует файл как UTF-8 с BOM, а при ошибке — как cp1251 (кодировка Excel на Windows)."""
    data = file.read()

    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251")


def detect_delimiter(text: str) -> str:
    """Определяет разделитель CSV по строке заголовка."""
    header = text.splitlines()[0]

    return ";" if header.count(";") > header.count(",") else ","


def make_post(data: dict) -> PostCreate:
    """Собирает пост из строки файла."""
    created_date = data["created_date"]

    if isinstance(created_date, str):
        created_date = datetime.strptime(
            created_date,
            "%d.%m.%Y %H:%M"
        )

    rubrics = data["rubrics"]

    if isinstance(rubrics, str):
        rubrics = ast.literal_eval(rubrics)

    return PostCreate(
        text=data["text"],
        created_date=created_date,
        rubrics=rubrics,
    )
